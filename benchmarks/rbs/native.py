"""Portable primary-source Tan TSS contexts; no scoring or translation labels."""

import argparse
import gzip
import json
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import unquote
from urllib.request import urlopen

from benchmarks.rbs.datasets import digest, write_json

MANIFEST = Path(__file__).with_name("native-sources.json")
REPLICON_MAP = {"CP006471": "NZ_CP006471.1", "CP006472": "NZ_CP006472.1"}
COMPLEMENT = str.maketrans("ACGT", "TGCA")
S1_HEADERS = ["TSS_ID", "Chromosome", "Type", "Strand", "TSS_position"]
S2_HEADERS = ["Type", "Position", "Locus_tag", "Product", "Distances to start codon"]


def manifest():
    """Read the versioned source ledger."""
    return json.loads(MANIFEST.read_text())


def checked_source(root, identity):
    """Return exact pinned bytes; changed caches are never silently replaced."""
    source = next(row for row in manifest()["sources"] if row["id"] == identity)
    data = (Path(root) / source["name"]).read_bytes()
    if len(data) != source["bytes"] or digest(data) != source["sha256"]:
        raise ValueError(f"Native source identity mismatch: {identity}")
    return data


def fetch_sources(root):
    """Fetch bounded declared files outside Git and verify bytes before writing."""
    root = Path(root).resolve()
    repository = Path(__file__).resolve().parents[2]
    if root == repository or repository in root.parents:
        raise ValueError("Raw source downloads must stay outside the repository")
    root.mkdir(parents=True, exist_ok=True)
    for source in manifest()["sources"]:
        path = root / source["name"]
        if not path.exists():
            with urlopen(source["url"], timeout=30) as response:
                data = response.read(source["bytes"] + 1)
            if len(data) != source["bytes"] or digest(data) != source["sha256"]:
                raise ValueError(f"Downloaded source identity mismatch: {source['id']}")
            path.write_bytes(data)
            path.chmod(0o600)
        checked_source(root, source["id"])
    return {"n_verified": len(manifest()["sources"])}


def fasta_records(text, alphabet="ACGT"):
    """Read exact versioned FASTA IDs in a declared alphabet; reject duplicate IDs."""
    records, key = {}, None
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith(">"):
            tokens = line[1:].split()
            if not tokens:
                raise ValueError("Empty FASTA identifier")
            key = tokens[0]
            if key in records:
                raise ValueError("Duplicate FASTA identifier")
            records[key] = ""
        else:
            if key is None or set(line.strip().upper()) - set(alphabet):
                raise ValueError("Invalid or unlabelled sequence")
            records[key] += line.strip().upper()
    if not records or any(not sequence for sequence in records.values()):
        raise ValueError("Empty genomic sequence")
    return records


def gff_features(text, sequences):
    """Preserve GFF's unwrapped circular intervals and transcript-ordered segments."""
    features = []
    circular = set()
    rows = []
    for line in text.splitlines():
        if not line or line.startswith("#"):
            continue
        fields = line.split("\t")
        if len(fields) != 9:
            raise ValueError("GFF must have nine columns")
        if fields[2] not in {"CDS", "rRNA", "region", "gene", "pseudogene"}:
            continue
        attrs = dict(item.split("=", 1) for item in fields[8].split(";") if item)
        attrs = {key: unquote(value) for key, value in attrs.items()}
        rows.append((fields, attrs))
        if fields[2] == "region" and attrs.get("Is_circular") == "true":
            circular.add(fields[0])
    for fields, attrs in rows:
        lower, upper = int(fields[3]) - 1, int(fields[4])
        if fields[0] not in sequences or fields[6] not in {"+", "-"}:
            raise ValueError("Invalid GFF sequence, strand or interval")
        length = len(sequences[fields[0]])
        if not 0 <= lower < length or not lower < upper <= lower + length:
            raise ValueError("Invalid GFF sequence, strand or interval")
        segments = [[lower, min(upper, length)]]
        if upper > length:
            if fields[0] not in circular or fields[2] == "region":
                raise ValueError("Origin-spanning GFF feature lacks circular topology")
            segments.append([0, upper - length])
        if fields[6] == "-":
            segments.reverse()
        features.append(
            {
                "replicon": fields[0],
                "type": fields[2],
                "strand": fields[6],
                "interval": [lower, upper],
                "attributes": attrs,
                "phase": fields[7],
                "segments": segments,
                "start_zero_based": (lower if fields[6] == "+" else upper - 1) % length,
            }
        )
    return features


def rrna_tail_evidence(sequences, features):
    """Return annotated 16S terminal nine bases, not a mature-RNA processing claim."""
    result = []
    for feature in features:
        if (
            feature["type"] != "rRNA"
            or feature["attributes"].get("product") != "16S ribosomal RNA"
        ):
            continue
        parts = [
            sequences[feature["replicon"]][lower:upper]
            for lower, upper in feature["segments"]
        ]
        if feature["strand"] == "-":
            parts = [part.translate(COMPLEMENT)[::-1] for part in parts]
        sequence = "".join(parts)
        if len(sequence) < 9:
            raise ValueError("16S annotation is shorter than the requested tail")
        result.append(
            {
                "locus_tag": feature["attributes"]["locus_tag"],
                "replicon": feature["replicon"],
                "strand": feature["strand"],
                "rrna_interval": feature["interval"],
                "rrna_sha256": digest(sequence),
                "anti_sd": sequence[-9:].replace("T", "U"),
                "evidence": (
                    "terminal 9 nt of versioned annotated 16S; processing unmeasured"
                ),
            }
        )
    return result


def tan_contexts(s1, s2, sequences, features):
    """Join primary TSS rows to exact current CDS IDs, preserving every source row."""
    if s1[1][:5] != S1_HEADERS or s2[1][:5] != S2_HEADERS:
        raise ValueError("Unexpected Tan workbook schema")
    genes, proposals = defaultdict(list), defaultdict(list)
    for feature in features:
        locus = feature["attributes"].get("locus_tag")
        if feature["type"] == "CDS" and locus:
            genes[locus].append(feature)
    notes = []
    for index, row in enumerate(s2[2:], start=3):
        if row[0] in {"gTSS", "iTSS"}:
            proposals[(row[0], row[1], row[2])].append(
                {
                    "source_row": index,
                    "need_reannotation": row[5],
                    "leaderless": row[6],
                    "misannotation_percent": row[7],
                    "leaderless_percent": row[8],
                    "homology_hit_count": row[9],
                    "evidence_type": "homology-based proposal",
                }
            )
        elif any(value is not None for value in row):
            notes.append({"source_row": index, "text": row[0]})
    result, seen = [], set()
    for index, row in enumerate(s1[2:], start=3):
        identity, chromosome, kind, strand, position = row[:5]
        if (
            identity in seen
            or kind not in {"gTSS", "aTSS", "iTSS", "nTSS"}
            or strand not in {"+", "-"}
            or chromosome not in REPLICON_MAP
            or type(position) is not int
        ):
            raise ValueError("Invalid or duplicate primary TSS record")
        seen.add(identity)
        replicon = REPLICON_MAP[chromosome]
        sequence, point, locus = sequences[replicon], position - 1, row[10]
        if not 0 <= point < len(sequence) or identity != f"{kind}{strand}{position}":
            raise ValueError("TSS identifier/coordinate mismatch")
        base = sequence[point]
        if strand == "-":
            base = base.translate(COMPLEMENT)
        record = {
            "id": "tan2018:" + identity,
            "study": "Tan2018",
            "source_sheet": "Table S1",
            "source_row": index,
            "source_locus_tag": locus,
            "genome_id": replicon,
            "replicon": replicon,
            "strand": strand,
            "tss_zero_based": point,
            "original_tss_one_based": position,
            "tss_type": kind,
            "transcript_start_evidence": "dRNA-seq TSS",
            "translation_start_evidence": "annotation",
            "source_start_distance_nt": row[9],
            "source_start_zero_based": None,
            "cds_start_zero_based": None,
            "leader_length_nt": None,
            "source_current_start_delta_nt": None,
            "origin_crossing": False,
            "upstream_limit_nt": 100,
            "downstream_length_nt": 100,
            "mapping_status": "non_gene_tss",
            "warnings": [],
            "source_flags": proposals.get((kind, position, locus), []),
            "source_tss_nucleotide": row[8],
            "reference_tss_nucleotide": base,
            "conditions": {
                name: list(row[a : a + 2])
                for name, a in [
                    ("control", 20),
                    ("dark", 22),
                    ("high_light", 24),
                    ("high_temperature", 26),
                ]
            },
            "source_quality": {s1[1][i]: row[i] for i in range(5, 8)},
        }
        if base != row[8]:
            record["warnings"].append("source_tss_nucleotide_disagrees_with_reference")
        if kind == "gTSS":
            if type(row[9]) is not int or row[9] < 0:
                raise ValueError("gTSS source distance must be a nonnegative integer")
            sign = 1 if strand == "+" else -1
            historical = (point + sign * row[9]) % len(sequence)
            record["source_start_zero_based"] = historical
            matches = genes.get(locus, [])
            if not matches:
                record["mapping_status"] = "source_locus_absent_from_current_cds"
            elif len(matches) != 1:
                record["mapping_status"] = "multiple_current_cds_segments_or_models"
            elif matches[0]["replicon"] != replicon or matches[0]["strand"] != strand:
                record["mapping_status"] = "current_replicon_or_strand_disagreement"
            else:
                feature = matches[0]
                current = feature["start_zero_based"]
                delta = sign * (current - historical)
                # Choose the local annotation shift, not a whole-genome rotation.
                delta = (delta + len(sequence) // 2) % len(sequence) - len(
                    sequence
                ) // 2
                leader = row[9] + delta
                record.update(
                    cds_start_zero_based=current,
                    leader_length_nt=leader,
                    source_current_start_delta_nt=delta,
                    current_cds_interval=feature["interval"],
                    current_cds_segments=feature["segments"],
                    current_cds_attributes=feature["attributes"],
                    mapping_status="matched",
                )
                if delta:
                    record["warnings"].append("source_current_start_disagreement")
                if leader < 0:
                    record["mapping_status"] = "tss_inside_current_cds"
                elif leader == 0:
                    record["warnings"].append("leaderless_current_annotation")
                else:
                    record["origin_crossing"] = sign * (current - point) != leader
                if feature["attributes"].get("pseudo") == "true":
                    record["mapping_status"] = "current_pseudogene"
        result.append(record)
    return result, {
        "n_s1": len(result),
        "types": dict(Counter(r["tss_type"] for r in result)),
        "mapping_status": dict(Counter(r["mapping_status"] for r in result)),
        "n_s2_proposals": sum(map(len, proposals.values())),
        "s2_notes": notes,
        "n_source_current_start_disagreements": sum(
            bool(r["source_current_start_delta_nt"]) for r in result
        ),
        "n_tss_nucleotide_disagreements": sum(
            r["source_tss_nucleotide"] != r["reference_tss_nucleotide"] for r in result
        ),
    }


def prepare_tan(root, output):
    """Prepare a genome-once source envelope using optional spreadsheet ingestion."""
    import io

    import openpyxl

    source = manifest()
    sequences = fasta_records(
        gzip.decompress(checked_source(root, "utex-genomic.fna.gz")).decode()
    )
    if {key: len(value) for key, value in sequences.items()} != source["utex_identity"][
        "replicons"
    ]:
        raise ValueError("UTEX replicon identity or length mismatch")
    features = gff_features(
        gzip.decompress(checked_source(root, "utex-genomic.gff.gz")).decode(), sequences
    )
    workbook_bytes = checked_source(root, "tan2018.xlsx")
    workbook = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes), read_only=True, data_only=True
    )
    s1 = [
        list(row)
        for row in workbook["Table S1"].iter_rows(max_col=30, values_only=True)
    ]
    s2 = [
        list(row)
        for row in workbook["Table S2"].iter_rows(max_col=21, values_only=True)
    ]
    workbook.close()
    records, audit = tan_contexts(s1, s2, sequences, features)
    tails = rrna_tail_evidence(sequences, features)
    asds = {row["anti_sd"] for row in tails}
    if len(asds) != 1:
        raise ValueError("No unique annotated 16S tail scenario")
    workbook_hash = digest(workbook_bytes)
    for row in records:
        row.update(
            host="Synechococcus elongatus UTEX 2973",
            anti_sd=next(iter(asds)),
            anti_sd_provenance=tails,
            source_file_sha256=workbook_hash,
        )
    regions = {row["replicon"]: row for row in features if row["type"] == "region"}
    if set(regions) != set(sequences) or any(
        row["attributes"].get("Is_circular") != "true" for row in regions.values()
    ):
        raise ValueError("UTEX circular topology is not established by this annotation")
    genomes = {
        key: {
            "sequence": value,
            "sequence_sha256": digest(value),
            "accession": key,
            "assembly": source["utex_identity"]["assembly"],
            "annotation": source["utex_identity"]["annotation"],
            "topology": "circular",
        }
        for key, value in sequences.items()
    }
    envelope = {
        "schema_version": "native-rbs-evidence-1",
        "genomes": genomes,
        "records": records,
        "audit": audit,
        "licence": "CC-BY-4.0 attribution retained; NCBI sequence facts",
        "attribution": "Tan et al. (2018), doi:10.1186/s13068-018-1215-8; NCBI RefSeq",
        "conditions_note": (
            "Two cultures/condition; high-light duration conflicts (30 min vs 2 h); "
            "high-temperature light level conflicts. Counts are not isoform weights "
            "or translation measurements."
        ),
    }
    write_json(output, envelope)
    return audit


def main(argv=None):
    """Fetch pinned primary files or produce a Tan context envelope."""
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    fetch = commands.add_parser("fetch")
    fetch.add_argument("root", type=Path)
    prepare = commands.add_parser("prepare-tan")
    prepare.add_argument("root", type=Path)
    prepare.add_argument("output", type=Path)
    args = parser.parse_args(argv)
    result = (
        fetch_sources(args.root)
        if args.command == "fetch"
        else prepare_tan(args.root, args.output)
    )
    print(json.dumps(result, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
